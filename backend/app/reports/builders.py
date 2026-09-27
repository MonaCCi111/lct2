"""Генерация отчётов: Excel (openpyxl) и PDF (reportlab). BACKEND_SANYA_GUIDE §5, INTEGRATION_SPEC §3.6."""
from __future__ import annotations

import io
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.timeutil import iso_msk, now_msk

FONT_DIR = Path(__file__).resolve().parent / "fonts"
BRAND_DARK = "1C1D22"
BRAND_ACCENT = "FF0053"
BRAND_SECONDARY = "8A83D1"
RISK_LABEL = {"critical": "Критический", "high": "Высокий", "medium": "Средний", "low": "Низкий", None: "—"}
STATUS_LABEL = {"draft": "Черновик", "approved": "Утверждён", "rejected": "Отклонён", "completed": "Выполнен"}

_fonts_ready = False


def _register_fonts() -> tuple[str, str]:
    global _fonts_ready
    regular, bold = "DejaVuSans", "DejaVuSans-Bold"
    if not _fonts_ready:
        candidates = [FONT_DIR, Path("/usr/share/fonts/truetype/dejavu")]
        for d in candidates:
            if (d / "DejaVuSans.ttf").exists():
                pdfmetrics.registerFont(TTFont(regular, str(d / "DejaVuSans.ttf")))
                b = d / "DejaVuSans-Bold.ttf"
                pdfmetrics.registerFont(TTFont(bold, str(b if b.exists() else d / "DejaVuSans.ttf")))
                _fonts_ready = True
                break
        if not _fonts_ready:  # без кириллического шрифта - стандартный Helvetica (кириллица не отобразится)
            regular, bold = "Helvetica", "Helvetica-Bold"
    return regular, bold


def _fmt_dt(v: datetime | str | None) -> str:
    if v is None:
        return ""
    if isinstance(v, str):
        return v.replace("T", " ")[:19]
    s = iso_msk(v) or ""
    return s.replace("T", " ")[:19]


# ------------------------------------------------------------------ XLSX
def build_incidents_xlsx(rows: list[dict], generated_at: datetime | None = None) -> bytes:
    """rows: словари с ключами ticket_id, created_at, object_name, piket, channel_id, sensor_name, sensor_type,
    failure_probability, risk_level, primary_cause, recommendation, status, assignee, prediction_id."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Реестр прогнозов и нарядов"
    generated_at = generated_at or now_msk()

    ws["A1"] = "АО «Москоллектор» — реестр предиктивных инцидентов и нарядов на ТО"
    ws["A1"].font = Font(bold=True, size=13, color=BRAND_ACCENT)
    ws["A2"] = f"Сформировано: {_fmt_dt(generated_at)} (МСК). Сервис прогнозирования инцидентов, команда Dolos."
    ws["A2"].font = Font(italic=True, size=9, color="666666")

    headers = [
        ("ID наряда", 16), ("Дата создания", 20), ("Объект", 26), ("Пикет", 10), ("Канал", 10), ("Датчик", 28),
        ("Тип датчика", 22), ("Вероятность отказа", 14), ("Уровень риска", 14), ("Причина", 38), ("Рекомендация", 40),
        ("Статус", 14), ("Бригада", 18), ("ID прогноза", 22),
    ]
    header_row = 4
    fill = PatternFill("solid", fgColor=BRAND_DARK)
    thin = Side(style="thin", color="BBBBBB")
    for col, (title, width) in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=title)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[header_row].height = 30

    risk_fill = {
        "critical": PatternFill("solid", fgColor="F8D7DA"),
        "high": PatternFill("solid", fgColor="FDE2C8"),
        "medium": PatternFill("solid", fgColor="FFF3CD"),
        "low": PatternFill("solid", fgColor="D1F2E4"),
    }
    for i, r in enumerate(rows, start=header_row + 1):
        values = [
            r.get("ticket_id") or "—", _fmt_dt(r.get("created_at")), r.get("object_name") or "", r.get("piket") or "",
            r.get("channel_id"), r.get("sensor_name") or "", r.get("sensor_type") or "",
            r.get("failure_probability"), RISK_LABEL.get(r.get("risk_level"), "—"), r.get("primary_cause") or "",
            r.get("recommendation") or "", STATUS_LABEL.get(r.get("status") or "", r.get("status") or "—"),
            r.get("assignee") or "", r.get("prediction_id") or "",
        ]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=i, column=col, value=v)
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            c.alignment = Alignment(vertical="top", wrap_text=col in (10, 11))
        if isinstance(values[7], (int, float)):
            ws.cell(row=i, column=8).number_format = "0%"
        if r.get("risk_level") in risk_fill:
            ws.cell(row=i, column=9).fill = risk_fill[r["risk_level"]]
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{max(header_row, header_row + len(rows))}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ PDF
def build_summary_pdf(summary: dict) -> bytes:
    """summary: generated_at, model_version, channels{...}, predictions{...}, objects{...}, tickets{...},
    collectors: [{name, risk_level, active, critical, high, channels_total}], critical_items: [{...}],
    prevented_incidents, reliability_index, weather{...} (опционально)."""
    regular, bold = _register_fonts()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
                            title="Справка о надёжности коллекторов", author="Dolos / Москоллектор")
    h1 = ParagraphStyle("h1", fontName=bold, fontSize=16, leading=20, textColor=colors.HexColor("#" + BRAND_ACCENT), spaceAfter=4)
    h2 = ParagraphStyle("h2", fontName=bold, fontSize=12, leading=15, textColor=colors.HexColor("#" + BRAND_DARK), spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("body", fontName=regular, fontSize=9.5, leading=13)
    small = ParagraphStyle("small", fontName=regular, fontSize=8, leading=10, textColor=colors.HexColor("#666666"))
    cell = ParagraphStyle("cell", fontName=regular, fontSize=8.5, leading=10.5)

    story = []
    # "Логотип": текстовый бренд-блок (векторного логотипа Москоллектора в репозитории нет)
    logo = Table([[Paragraph("<b>МОСКОЛЛЕКТОР</b>", ParagraphStyle("logo", fontName=bold, fontSize=13, textColor=colors.white))]],
                 colWidths=[60 * mm], rowHeights=[12 * mm])
    logo.hAlign = "LEFT"
    logo.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#" + BRAND_DARK)),
                              ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 8)]))
    story += [logo, Spacer(1, 6)]
    story.append(Paragraph("Справка о надёжности инженерных коллекторов", h1))
    story.append(Paragraph(f"Дата выгрузки: {_fmt_dt(summary.get('generated_at'))} (МСК). Модель: {summary.get('model_version', '')}. "
                           f"Сформировано сервисом прогнозирования инцидентов (команда Dolos, ЛЦТ 2026).", small))

    story.append(Paragraph("Ключевые метрики", h2))
    ch, pr, ob, tk = summary.get("channels", {}), summary.get("predictions", {}), summary.get("objects", {}), summary.get("tickets", {})
    kpi = [
        ["Каналов под контролем", f"{ch.get('total', 0)}", "Покрытие ML-моделями", f"{ch.get('coverage_percent', 0)}%"],
        ["Активных прогнозов", f"{pr.get('active', 0)}", "Критических / высоких", f"{pr.get('critical', 0)} / {pr.get('high', 0)}"],
        ["Объектов в зоне риска", f"{ob.get('affected', 0)} из {ob.get('total', 0)}", "Индекс надёжности коллекторов", f"{summary.get('reliability_index', '—')}"],
        ["Предотвращённых аварий (нарядов выполнено)", f"{summary.get('prevented_incidents', 0)}", "Нарядов в работе (черновик / утверждён)", f"{tk.get('draft', 0)} / {tk.get('approved', 0)}"],
    ]
    t = Table([[Paragraph(str(x), cell) for x in row] for row in kpi], colWidths=[58 * mm, 28 * mm, 58 * mm, 28 * mm])
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F3F3F6")), ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#F3F3F6")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(t)

    w = summary.get("weather")
    if w:
        story.append(Paragraph(
            f"Метеоусловия Москвы на момент выгрузки: {w.get('temperature_c')} °C, влажность {w.get('relative_humidity')}%, "
            f"давление {w.get('surface_pressure_hpa')} гПа, осадки {w.get('precipitation_mm')} мм ({w.get('source')}).", small))

    story.append(Paragraph("Сводка по коллекторам", h2))
    head = ["Объект", "Уровень риска", "Активных прогнозов", "Критических", "Высоких", "Каналов"]
    rows = [[Paragraph(f"<b>{h}</b>", cell) for h in head]]
    for c in summary.get("collectors", []):
        rows.append([Paragraph(str(c.get("name", "")), cell), Paragraph(RISK_LABEL.get(c.get("risk_level"), "—"), cell),
                     str(c.get("active", 0)), str(c.get("critical", 0)), str(c.get("high", 0)), str(c.get("channels_total", 0))])
    t = Table(rows, colWidths=[62 * mm, 28 * mm, 30 * mm, 22 * mm, 18 * mm, 20 * mm], repeatRows=1)
    style = [("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
             ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E9E7F5")),
             ("ALIGN", (2, 1), (-1, -1), "CENTER"), ("FONTNAME", (0, 0), (-1, -1), regular), ("FONTSIZE", (0, 0), (-1, -1), 8.5)]
    for i, c in enumerate(summary.get("collectors", []), start=1):
        if c.get("risk_level") == "critical":
            style.append(("BACKGROUND", (1, i), (1, i), colors.HexColor("#F8D7DA")))
        elif c.get("risk_level") == "high":
            style.append(("BACKGROUND", (1, i), (1, i), colors.HexColor("#FDE2C8")))
    t.setStyle(TableStyle(style))
    story.append(t)

    story.append(Paragraph("Критические риски (требуется наряд на ТО)", h2))
    items = summary.get("critical_items", [])
    if not items:
        story.append(Paragraph("Критических рисков на момент выгрузки нет.", body))
    else:
        head = ["Объект / пикет", "Датчик (канал)", "Риск", "Причина", "Рекомендация", "Наряд"]
        rows = [[Paragraph(f"<b>{h}</b>", cell) for h in head]]
        for it in items[:40]:
            rows.append([
                Paragraph(f"{it.get('object_name', '')}<br/>{it.get('piket') or ''}", cell),
                Paragraph(f"{it.get('sensor_name', '')} ({it.get('channel_id', '')})", cell),
                Paragraph(f"{round(100 * (it.get('failure_probability') or 0))}%", cell),
                Paragraph(str(it.get("primary_cause") or ""), cell),
                Paragraph(str(it.get("recommendation") or ""), cell),
                Paragraph(f"{it.get('ticket_id') or '—'}<br/>{STATUS_LABEL.get(it.get('ticket_status') or '', '')}", cell),
            ])
        t = Table(rows, colWidths=[34 * mm, 34 * mm, 18 * mm, 40 * mm, 40 * mm, 20 * mm], repeatRows=1)
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
                               ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E9E7F5")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(t)
        if len(items) > 40:
            story.append(Paragraph(f"Показаны 40 из {len(items)} критических рисков; полный перечень - в Excel-выгрузке.", small))

    story.append(Spacer(1, 8))
    story.append(Paragraph("Вероятности рассчитаны на горизонте 24 ч по паттернам телеметрии СМВУ (дребезг контактов, доля тревог, "
                           "нестабильность показаний). Справка носит рекомендательный характер; решение о выезде принимает диспетчер ОДС.", small))
    doc.build(story)
    return buf.getvalue()
